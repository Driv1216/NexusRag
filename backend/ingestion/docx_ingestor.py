import os
from typing import List, Dict, Any, Tuple
import docx
from pptx import Presentation
from backend.ingestion.base import BaseIngestor
from backend.ingestion.chunking import HierarchicalChunker
from backend.models import DocumentChunk, ChunkMetadata

class DocxPptxIngestor(BaseIngestor):
    """Parses Microsoft Word (.docx) and PowerPoint (.pptx) documents."""

    def __init__(self):
        self.chunker = HierarchicalChunker()

    def ingest(self, file_path: str, filename: str) -> Tuple[List[DocumentChunk], Dict[str, Any]]:
        ext = os.path.splitext(filename)[1].lower()
        if ext == ".docx":
            return self._ingest_docx(file_path, filename)
        elif ext in [".pptx", ".ppt"]:
            return self._ingest_pptx(file_path, filename)
        else:
            raise ValueError(f"Unsupported office document format: {ext}")

    def _ingest_docx(self, file_path: str, filename: str) -> Tuple[List[DocumentChunk], Dict[str, Any]]:
        doc = docx.Document(file_path)
        full_text_blocks = []
        tables_count = 0

        # Paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                full_text_blocks.append(text)

        combined_text = "\n\n".join(full_text_blocks)
        chunks = self.chunker.chunk_document_section(
            text=combined_text,
            source_file=filename,
            page_number=1,
            modality="text"
        )

        # Tables
        clean_name = os.path.splitext(filename)[0]
        for t_idx, table in enumerate(doc.tables):
            rows_data = []
            for row in table.rows:
                row_cells = [c.text.replace("\n", " ").strip() for c in row.cells]
                rows_data.append(row_cells)
            
            if rows_data:
                tables_count += 1
                headers = rows_data[0]
                md_rows = [f"| {' | '.join(headers)} |", f"| {' | '.join(['---'] * len(headers))} |"]
                for r in rows_data[1:]:
                    md_rows.append(f"| {' | '.join(r[:len(headers)])} |")
                table_md = "\n".join(md_rows)
                
                tbl_id = f"table_{clean_name}_{t_idx+1}"
                chunks.append(DocumentChunk(
                    chunk_id=tbl_id,
                    content=f"[DOCX TABLE: {filename} Table {t_idx+1}]\n{table_md}",
                    metadata=ChunkMetadata(
                        chunk_id=tbl_id,
                        source_file=filename,
                        modality="table",
                        page_number=1,
                        table_html=table_md
                    )
                ))

        return chunks, {
            "total_pages": max(1, len(doc.paragraphs) // 10),
            "figures_extracted": 0,
            "tables_extracted": tables_count,
            "total_chunks": len(chunks)
        }

    def _ingest_pptx(self, file_path: str, filename: str) -> Tuple[List[DocumentChunk], Dict[str, Any]]:
        prs = Presentation(file_path)
        chunks: List[DocumentChunk] = []
        clean_name = os.path.splitext(filename)[0]
        total_slides = len(prs.slides)

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            slide_text_parts = []
            
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        txt = paragraph.text.strip()
                        if txt:
                            slide_text_parts.append(txt)

            slide_text = "\n".join(slide_text_parts)
            if slide_text:
                slide_chunks = self.chunker.chunk_document_section(
                    text=f"[SLIDE {slide_num}]\n" + slide_text,
                    source_file=filename,
                    page_number=slide_num,
                    modality="text",
                    extra_meta={"slide_number": slide_num}
                )
                chunks.extend(slide_chunks)

        return chunks, {
            "total_pages": total_slides,
            "figures_extracted": 0,
            "tables_extracted": 0,
            "total_chunks": len(chunks)
        }
